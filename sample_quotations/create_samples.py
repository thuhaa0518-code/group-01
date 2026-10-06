import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

os.makedirs('sample_quotations', exist_ok=True)

# -------------------------------------------------------------
# 1. Excel Quotation: HaNoiTech (Thiết bị CNTT)
# -------------------------------------------------------------
wb1 = openpyxl.Workbook()
ws1 = wb1.active
ws1.title = "Bao Gia CNTT"

# Header
ws1['A1'] = "CÔNG TY TNHH THIẾT BỊ CÔNG NGHỆ HÀ NỘI"
ws1['A1'].font = Font(name="Arial", size=14, bold=True, color="1E40AF")
ws1['A2'] = "Mã số thuế: 0102345678 | Hotline: 024-3852-1122 | Email: hung.nv@hanoitech.com.vn"
ws1['A2'].font = Font(name="Arial", size=9, italic=True, color="64748B")
ws1['A3'] = "Địa chỉ: Số 102 Đường Giải Phóng, Q. Đống Đa, Hà Nội"
ws1['A3'].font = Font(name="Arial", size=9, italic=True, color="64748B")

ws1['A5'] = "BẢNG BÁO GIÁ THIẾT BỊ CÔNG NGHỆ THÔNG TIN"
ws1['A5'].font = Font(name="Arial", size=15, bold=True, color="0F172A")

headers = ["STT", "Tên mặt hàng / Quy cách kỹ thuật", "ĐVT", "Số lượng", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)"]
ws1.append([])
ws1.append(headers)

# Items
items1 = [
    (1, "Laptop văn phòng 14 inch Core i7-1355U, 32GB RAM, 1TB SSD", "Chiếc", 5, 23500000),
    (2, "Màn hình 27 inch 2K IPS 75Hz (HDMI / USB-C)", "Chiếc", 4, 6200000),
    (3, "Bộ bàn phím + chuột không dây Bluetooth", "Bộ", 3, 1150000),
]

for row in items1:
    ws1.append([row[0], row[1], row[2], row[3], row[4], row[3] * row[4]])

ws1.append([])
ws1.append(["", "CỘNG TIỀN HÀNG (Chưa VAT):", "", "", "", 145750000])
ws1.append(["", "Thuế GTGT VAT (10%):", "", "", "", 14575000])
ws1.append(["", "Phí vận chuyển & lắp đặt tận nơi:", "", "", "", 300000])
ws1.append(["", "TỔNG CỘNG THANH TOÁN:", "", "", "", 160625000])

ws1['A15'] = "ĐIỀU KIỆN THƯƠNG MẠI & BẢO HÀNH:"
ws1['A15'].font = Font(name="Arial", size=10, bold=True)
ws1['A16'] = "• Thời gian giao hàng: Trong vòng 3 ngày làm việc kể từ ngày phát hành PO."
ws1['A17'] = "• Thời hạn bảo hành: 24 tháng đối với Laptop & Màn hình, 12 tháng đối với Phụ kiện."
ws1['A18'] = "• Hiệu lực báo giá: 30 ngày kể từ ngày báo giá."

wb1.save('sample_quotations/BaoGia_Laptop_HaNoiTech.xlsx')

# -------------------------------------------------------------
# 2. Excel Quotation: MinhAnMedia (In ấn & Marketing)
# -------------------------------------------------------------
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "Bao Gia In An"

ws2['A1'] = "CÔNG TY TNHH TRUYỀN THÔNG & IN ẤN MINH AN"
ws2['A1'].font = Font(name="Arial", size=14, bold=True, color="047857")
ws2['A2'] = "Mã số thuế: 0105678901 | Email: an.pm@minhanmedia.vn | SĐT: 024-3512-8899"
ws2['A2'].font = Font(name="Arial", size=9, italic=True, color="64748B")

ws2['A4'] = "BÁO GIÁ IN ẤN ẤN PHẨM MARKETING QUÝ 4"
ws2['A4'].font = Font(name="Arial", size=14, bold=True)

headers2 = ["STT", "Hạng mục ấn phẩm / Quy cách", "ĐVT", "Số lượng", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)"]
ws2.append(headers2)

items2 = [
    (1, "Brochure A4 gấp 3 (Giấy C150, in 4 màu 2 mặt, cán mờ)", "Tờ", 2000, 8500),
    (2, "Standee cuốn 80x200cm (Bạt Hiflex, khung nhôm)", "Bộ", 10, 1400000),
    (3, "Backdrop sân khấu 6x3m (Khung sắt, bạt in UV)", "Bộ", 1, 4800000),
]

for row in items2:
    ws2.append([row[0], row[1], row[2], row[3], row[4], row[3] * row[4]])

ws2.append([])
ws2.append(["", "Tổng tiền trước thuế:", "", "", "", 35800000])
ws2.append(["", "Thuế GTGT VAT (10%):", "", "", "", 3580000])
ws2.append(["", "Tổng cộng thanh toán:", "", "", "", 39380000])

wb2.save('sample_quotations/BaoGia_InAn_MinhAnMedia.xlsx')

# -------------------------------------------------------------
# 3. Excel Quotation: TanPhat (Nội thất Văn phòng)
# -------------------------------------------------------------
wb3 = openpyxl.Workbook()
ws3 = wb3.active
ws3.title = "Bao Gia Noi That"

ws3['A1'] = "CÔNG TY CỔ PHẦN NỘI THẤT VĂN PHÒNG TÂN PHÁT"
ws3['A1'].font = Font(name="Arial", size=14, bold=True, color="B45309")
ws3['A2'] = "Mã số thuế: 0104321098 | Hotline: 024-3628-4433 | Email: phat.lt@tanphatfurniture.vn"
ws3['A2'].font = Font(name="Arial", size=9, italic=True)

ws3['A4'] = "BÁO GIÁ THIẾT BỊ NỘI THẤT CÔNG THÁI HỌC"
ws3['A4'].font = Font(name="Arial", size=14, bold=True)

headers3 = ["STT", "Tên sản phẩm / Quy cách", "ĐVT", "Số lượng", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)"]
ws3.append(headers3)

items3 = [
    (1, "Ghế công thái học lưng lưới TP-E08 (Tựa đầu, tay 4D, đệm trượt)", "Chiếc", 10, 4800000),
    (2, "Bàn nâng hạ chiều cao điện tử 140x70cm", "Bộ", 2, 7500000),
]

for row in items3:
    ws3.append([row[0], row[1], row[2], row[3], row[4], row[3] * row[4]])

ws3.append([])
ws3.append(["", "Cộng tiền hàng:", "", "", "", 63000000])
ws3.append(["", "Thuế VAT (10%):", "", "", "", 6300000])
ws3.append(["", "Phí vận chuyển & lắp đặt:", "", "", "", 500000])
ws3.append(["", "Tổng cộng thanh toán:", "", "", "", 69800000])

wb3.save('sample_quotations/BaoGia_GheCongThaiHoc_TanPhat.xlsx')

print("Created 3 Excel sample quotation files in sample_quotations/")
